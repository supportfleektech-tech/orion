"""Turn any uploaded file into something the model can actually reason about.

Images go to the model as base64 (when it is vision-capable); everything else is
converted to text. Each attachment reports how it was handled so the UI and the
model both know whether they are seeing pixels, extracted text, or a fallback
description -- ORION never pretends it understood a file it could not read.
"""

from __future__ import annotations

import base64
import io
import logging
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

log = logging.getLogger(__name__)

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}
AUDIO_SUFFIXES = {".mp3", ".wav", ".ogg", ".m4a", ".flac", ".opus", ".aac"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".webm", ".mkv", ".avi"}
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".yaml", ".yml", ".toml", ".ini",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css", ".sql", ".sh", ".log",
    ".rst", ".xml", ".java", ".go", ".rs", ".c", ".cpp", ".h", ".rb", ".php", ".env",
}

MAX_IMAGE_EDGE = 1280          # downscale before base64 to protect context + RAM
MAX_TEXT_CHARS = 20_000
MAX_IMAGE_BYTES = 12 * 1024 * 1024


@dataclass
class Attachment:
    """A processed user upload, ready for the model."""

    name: str
    media_type: str                       # image | audio | video | document | text | unknown
    mime: str | None = None
    size_bytes: int = 0
    text: str | None = None               # extracted text, if any
    image_b64: str | None = None          # for vision models
    handled_as: str = "unsupported"       # how we actually processed it
    note: str | None = None               # honest limitation, shown to user and model
    meta: dict[str, Any] = field(default_factory=dict)

    def to_public(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "media_type": self.media_type,
            "mime": self.mime,
            "size_bytes": self.size_bytes,
            "handled_as": self.handled_as,
            "note": self.note,
            "has_image": bool(self.image_b64),
            "text_preview": (self.text or "")[:400] or None,
            "meta": self.meta,
        }


def classify(name: str) -> tuple[str, str]:
    suffix = Path(name).suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        return "image", suffix
    if suffix in AUDIO_SUFFIXES:
        return "audio", suffix
    if suffix in VIDEO_SUFFIXES:
        return "video", suffix
    if suffix in {".pdf", ".docx", ".pptx", ".xlsx", ".xls"}:
        return "document", suffix
    if suffix in TEXT_SUFFIXES:
        return "text", suffix
    return "unknown", suffix


# ------------------------------------------------------------------ images
def process_image(data: bytes, name: str) -> Attachment:
    attachment = Attachment(name=name, media_type="image", mime=mimetypes.guess_type(name)[0], size_bytes=len(data))
    if len(data) > MAX_IMAGE_BYTES:
        attachment.note = f"Image is larger than {MAX_IMAGE_BYTES // 1024 // 1024} MB and was rejected."
        return attachment
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as image:
            image.load()
            attachment.meta = {"width": image.width, "height": image.height, "format": image.format}
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            if max(image.size) > MAX_IMAGE_EDGE:
                ratio = MAX_IMAGE_EDGE / max(image.size)
                image = image.resize((int(image.width * ratio), int(image.height * ratio)))
                attachment.meta["resized_to"] = list(image.size)
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=85)
            attachment.image_b64 = base64.b64encode(buffer.getvalue()).decode()
        attachment.handled_as = "image_for_vision_model"
    except Exception as exc:
        attachment.note = f"Could not decode this image: {exc}"
        log.warning("Image processing failed for %s: %s", name, exc)
    return attachment


# --------------------------------------------------------------- documents
def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    parts = []
    for index, slide in enumerate(deck.slides, start=1):
        parts.append(f"--- Slide {index} ---")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(max_row=200, values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


DOC_EXTRACTORS = {".pdf": _pdf_text, ".docx": _docx_text, ".pptx": _pptx_text, ".xlsx": _xlsx_text, ".xls": _xlsx_text}


def process_document(data: bytes, name: str, suffix: str) -> Attachment:
    attachment = Attachment(
        name=name, media_type="document", mime=mimetypes.guess_type(name)[0], size_bytes=len(data)
    )
    extractor = DOC_EXTRACTORS.get(suffix)
    if not extractor:
        attachment.note = f"No extractor for {suffix} files."
        return attachment
    try:
        text = (extractor(data) or "").strip()
        if not text:
            attachment.note = (
                "No selectable text found -- this is likely a scanned document. "
                "OCR is not enabled, so its contents could not be read."
            )
            attachment.handled_as = "empty_extraction"
            return attachment
        attachment.text = text[:MAX_TEXT_CHARS]
        attachment.handled_as = "text_extracted"
        attachment.meta = {"characters": len(text), "truncated": len(text) > MAX_TEXT_CHARS}
    except Exception as exc:
        attachment.note = f"Could not read this document: {exc}"
        log.warning("Document extraction failed for %s: %s", name, exc)
    return attachment


# -------------------------------------------------------------- transcripts
def process_audio(data: bytes, name: str) -> Attachment:
    """Transcribe audio when faster-whisper is installed; otherwise say so."""
    attachment = Attachment(
        name=name, media_type="audio", mime=mimetypes.guess_type(name)[0], size_bytes=len(data)
    )
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        attachment.note = (
            "Audio transcription is not enabled. Install it with "
            "`pip install -r backend/requirements-optional.txt` to transcribe voice input locally."
        )
        return attachment

    import tempfile

    try:
        from app.core.config import settings

        with tempfile.NamedTemporaryFile(suffix=Path(name).suffix, delete=True) as handle:
            handle.write(data)
            handle.flush()
            model = WhisperModel(settings.whisper_model, device="cpu", compute_type="int8")
            segments, info = model.transcribe(handle.name, beam_size=1)
            text = " ".join(segment.text.strip() for segment in segments).strip()
        if not text:
            attachment.note = "No speech detected in this audio."
            attachment.handled_as = "empty_transcription"
            return attachment
        attachment.text = text[:MAX_TEXT_CHARS]
        attachment.handled_as = "transcribed"
        attachment.meta = {"language": info.language, "duration_s": round(info.duration, 1)}
    except Exception as exc:
        attachment.note = f"Transcription failed: {exc}"
        log.warning("Transcription failed for %s: %s", name, exc)
    return attachment


def process_video(data: bytes, name: str) -> Attachment:
    attachment = Attachment(
        name=name, media_type="video", mime=mimetypes.guess_type(name)[0], size_bytes=len(data)
    )
    attachment.note = (
        "Video is accepted but not analysed frame-by-frame. Extract the audio track or a key frame "
        "and upload that for analysis."
    )
    return attachment


def process_text(data: bytes, name: str) -> Attachment:
    attachment = Attachment(
        name=name, media_type="text", mime=mimetypes.guess_type(name)[0], size_bytes=len(data)
    )
    text = data.decode("utf-8", errors="replace").strip()
    if not text:
        attachment.note = "File is empty."
        attachment.handled_as = "empty"
        return attachment
    attachment.text = text[:MAX_TEXT_CHARS]
    attachment.handled_as = "text_read"
    attachment.meta = {"characters": len(text), "truncated": len(text) > MAX_TEXT_CHARS}
    return attachment


def process_upload(data: bytes, name: str) -> Attachment:
    """Route an uploaded file to the right processor."""
    media_type, suffix = classify(name)
    if media_type == "image":
        return process_image(data, name)
    if media_type == "document":
        return process_document(data, name, suffix)
    if media_type == "audio":
        return process_audio(data, name)
    if media_type == "video":
        return process_video(data, name)
    if media_type == "text":
        return process_text(data, name)

    # Unknown: try UTF-8, else declare it binary rather than guessing.
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        return Attachment(
            name=name,
            media_type="unknown",
            size_bytes=len(data),
            note=f"Unsupported binary file type ({suffix or 'no extension'}); contents were not read.",
        )
    attachment = process_text(data, name)
    attachment.media_type = "text"
    return attachment


def build_user_content(prompt: str, attachments: list[Attachment], vision_capable: bool) -> Any:
    """Assemble an OpenAI-format user message from text + attachments."""
    text_parts: list[str] = [prompt] if prompt.strip() else []
    images: list[str] = []

    for attachment in attachments:
        if attachment.image_b64:
            if vision_capable:
                images.append(attachment.image_b64)
                text_parts.append(f"[Attached image: {attachment.name}]")
            else:
                dims = attachment.meta.get("width"), attachment.meta.get("height")
                text_parts.append(
                    f"[Attached image: {attachment.name} ({dims[0]}x{dims[1]}). "
                    "The active model has no vision capability, so this image was NOT analysed. "
                    "Tell the user to switch to a multimodal model such as qwen3.5:4b.]"
                )
        elif attachment.text:
            text_parts.append(
                f"--- Attached {attachment.media_type}: {attachment.name} "
                f"({attachment.handled_as}) ---\n{attachment.text}"
            )
        elif attachment.note:
            text_parts.append(f"[Attachment {attachment.name}: {attachment.note}]")

    combined = "\n\n".join(text_parts) or "(no text provided)"
    if not images:
        return combined

    content: list[dict[str, Any]] = [{"type": "text", "text": combined}]
    for image in images:
        content.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{image}"}})
    return content
