from __future__ import annotations

import hashlib
import mimetypes
from pathlib import Path

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models import Chunk, Document
from app.services.embeddings import cosine, embed_text, keyword_score
from app.services.reranker import rerank

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 180
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".json", ".csv", ".yaml", ".yml", ".toml", ".ini",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".sql", ".sh", ".log", ".rst",
}


def knowledge_root() -> Path:
    root = Path(settings.knowledge_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def read_source(path: Path) -> str:
    suffix = path.suffix.lower()

    # HTML is checked before the plain-text set, which also contains .html:
    # matching text first made the parser below unreachable and indexed pages
    # as raw markup, so a search for "script" hit every page on the site.
    if suffix in {".htm", ".html"}:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        # Script and style content is code, not prose; indexing it pollutes
        # retrieval with minified JavaScript.
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return soup.get_text(" ", strip=True)

    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document as DocxDocument

        return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
    raise ValueError(f"Unsupported file type: {suffix or path.name}")


def chunk_text(text_value: str) -> list[str]:
    text_value = " ".join((text_value or "").split())
    if not text_value:
        return []
    chunks: list[str] = []
    start = 0
    step = max(1, CHUNK_SIZE - CHUNK_OVERLAP)
    while start < len(text_value):
        chunks.append(text_value[start : start + CHUNK_SIZE])
        start += step
    return chunks


async def ingest_text(
    db: Session, name: str, content: str, source_path: str = "inline", meta: dict | None = None
) -> dict:
    digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
    existing = db.scalar(select(Document).where(Document.hash == digest))
    if existing:
        return {
            "document_id": existing.id,
            "name": existing.name,
            "chunks": existing.chunk_count,
            "status": "unchanged",
        }

    stale = db.scalars(select(Document).where(Document.path == source_path)).all()
    for doc in stale:
        db.execute(delete(Chunk).where(Chunk.document_id == doc.id))
        db.delete(doc)

    document = Document(
        name=name,
        path=source_path,
        mime_type=mimetypes.guess_type(name)[0],
        hash=digest,
        size_bytes=len(content.encode("utf-8")),
        meta=meta or {},
    )
    db.add(document)
    db.flush()

    pieces = chunk_text(content)
    for index, piece in enumerate(pieces):
        db.add(
            Chunk(
                document_id=document.id,
                chunk_index=index,
                content=piece,
                embedding=await embed_text(piece),
                meta={"name": name},
            )
        )
    document.chunk_count = len(pieces)
    db.commit()
    return {"document_id": document.id, "name": name, "chunks": len(pieces), "status": "ingested"}


async def ingest_path(db: Session, raw_path: str) -> dict:
    root = knowledge_root()
    candidate = Path(raw_path)
    target = candidate if candidate.is_absolute() else (root / candidate)
    target = target.resolve()
    if root != target and root not in target.parents:
        raise PermissionError("Path escapes the knowledge directory")

    if target.is_dir():
        results = []
        for file in sorted(target.rglob("*")):
            if file.is_file() and (file.suffix.lower() in TEXT_SUFFIXES or file.suffix.lower() in {".pdf", ".docx"}):
                try:
                    results.append(await ingest_path(db, str(file)))
                except Exception as exc:
                    results.append({"name": file.name, "status": "error", "error": str(exc)})
        return {"status": "batch", "count": len(results), "results": results}

    if not target.is_file():
        raise FileNotFoundError(raw_path)

    content = read_source(target)
    return await ingest_text(db, target.name, content, str(target))


async def search_chunks(db: Session, query: str, limit: int | None = None) -> list[dict]:
    """Hybrid first-stage retrieval followed by a reranking pass.

    Stage one optimises recall: take a wider candidate set than we need. Stage
    two reorders it for precision, so the chunk that actually answers the
    question lands at the top of the context window rather than fourth.
    """
    limit = limit or settings.max_context_chunks
    rows = db.scalars(select(Chunk).limit(5000)).all()
    if not rows:
        return []
    qvec = await embed_text(query)
    scored = []
    for chunk in rows:
        score = 0.65 * cosine(qvec, chunk.embedding or []) + 0.35 * keyword_score(query, chunk.content)
        scored.append((score, chunk))
    scored.sort(key=lambda x: x[0], reverse=True)

    pool = max(limit, settings.rerank_candidates) if settings.rerank_enabled else limit
    candidates = [
        {
            "id": c.id,
            "document_id": c.document_id,
            "name": (c.meta or {}).get("name"),
            "chunk_index": c.chunk_index,
            "content": c.content,
            "score": round(s, 4),
        }
        for s, c in scored[:pool]
        if s > 0.02
    ]
    return rerank(query, candidates, limit=limit)


def list_documents(db: Session, limit: int = 200, offset: int = 0) -> list[dict]:
    """Newest documents first.

    Bounded because this backs a UI list that grows with every ingest: a
    knowledge base of a few thousand files would otherwise serialise the whole
    table on every page load.
    """
    docs = db.scalars(
        select(Document).order_by(Document.created_at.desc()).limit(limit).offset(offset)
    ).all()
    return [
        {
            "id": d.id,
            "name": d.name,
            "path": d.path,
            "mime_type": d.mime_type,
            "size_bytes": d.size_bytes,
            "chunks": d.chunk_count,
            "created_at": d.created_at.isoformat() if d.created_at else None,
        }
        for d in docs
    ]


def delete_document(db: Session, document_id: str) -> bool:
    doc = db.get(Document, document_id)
    if not doc:
        return False
    db.execute(delete(Chunk).where(Chunk.document_id == doc.id))
    db.delete(doc)
    db.commit()
    return True
