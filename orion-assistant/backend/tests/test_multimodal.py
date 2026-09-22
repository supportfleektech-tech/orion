"""Attachment processing: every input format resolves to something usable or an honest note."""

from __future__ import annotations

import io

import pytest

from app.services.multimodal import (
    Attachment,
    build_user_content,
    classify,
    process_upload,
)


def png_bytes(width: int = 40, height: int = 30, color: str = "red") -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.parametrize(
    "name,expected",
    [
        ("photo.PNG", "image"),
        ("clip.mp3", "audio"),
        ("movie.mov", "video"),
        ("report.pdf", "document"),
        ("deck.pptx", "document"),
        ("sheet.xlsx", "document"),
        ("main.py", "text"),
        ("notes.md", "text"),
        ("archive.zip", "unknown"),
    ],
)
def test_classify(name, expected):
    assert classify(name)[0] == expected


def test_text_file_is_read():
    attachment = process_upload(b"hello orion", "notes.txt")
    assert attachment.media_type == "text"
    assert attachment.handled_as == "text_read"
    assert attachment.text == "hello orion"


def test_empty_text_file_is_flagged():
    attachment = process_upload(b"   ", "empty.txt")
    assert attachment.handled_as == "empty"
    assert attachment.note


def test_image_is_encoded_for_vision():
    attachment = process_upload(png_bytes(), "shot.png")
    assert attachment.media_type == "image"
    assert attachment.handled_as == "image_for_vision_model"
    assert attachment.image_b64
    assert attachment.meta["width"] == 40


def test_large_image_is_downscaled():
    attachment = process_upload(png_bytes(2000, 1500), "big.png")
    assert attachment.image_b64
    assert max(attachment.meta["resized_to"]) == 1280


def test_corrupt_image_reports_honestly():
    attachment = process_upload(b"not really a png", "broken.png")
    assert attachment.image_b64 is None
    assert "could not decode" in (attachment.note or "").lower()


def test_binary_unknown_type_is_not_guessed():
    attachment = process_upload(b"\x00\x01\x02\xff\xfe", "thing.bin")
    assert attachment.media_type == "unknown"
    assert "unsupported" in (attachment.note or "").lower()


def test_xlsx_text_extraction():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["region", "revenue"])
    workbook.active.append(["Nairobi", 4200])
    buffer = io.BytesIO()
    workbook.save(buffer)

    attachment = process_upload(buffer.getvalue(), "sales.xlsx")
    assert attachment.handled_as == "text_extracted"
    assert "Nairobi" in attachment.text
    assert "4200" in attachment.text


def test_docx_text_extraction():
    from docx import Document

    document = Document()
    document.add_paragraph("Quarterly review for ORION")
    buffer = io.BytesIO()
    document.save(buffer)

    attachment = process_upload(buffer.getvalue(), "review.docx")
    assert attachment.handled_as == "text_extracted"
    assert "Quarterly review" in attachment.text


def test_video_is_accepted_but_declared_unanalysed():
    attachment = process_upload(b"\x00" * 100, "clip.mp4")
    assert attachment.media_type == "video"
    assert "not analysed" in (attachment.note or "")


def test_audio_without_whisper_says_so():
    attachment = process_upload(b"\x00" * 100, "voice.wav")
    assert attachment.media_type == "audio"
    # Either it transcribed (whisper installed) or it told us how to enable it.
    assert attachment.text or "transcription" in (attachment.note or "").lower()


def test_build_content_sends_images_to_vision_model():
    attachment = process_upload(png_bytes(), "shot.png")
    content = build_user_content("what is this?", [attachment], vision_capable=True)
    assert isinstance(content, list)
    assert content[0]["type"] == "text"
    assert content[1]["type"] == "image_url"
    assert content[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")


def test_build_content_warns_when_model_is_blind():
    attachment = process_upload(png_bytes(), "shot.png")
    content = build_user_content("what is this?", [attachment], vision_capable=False)
    assert isinstance(content, str)
    assert "no vision capability" in content
    assert "qwen3.5:4b" in content


def test_build_content_inlines_extracted_text():
    attachment = process_upload(b"budget line items", "budget.txt")
    content = build_user_content("summarise", [attachment], vision_capable=False)
    assert "budget line items" in content
    assert "budget.txt" in content


def test_build_content_surfaces_failure_notes():
    broken = Attachment(name="x.bin", media_type="unknown", note="Unsupported binary file type")
    content = build_user_content("hi", [broken], vision_capable=False)
    assert "Unsupported binary file type" in content


def test_public_view_never_leaks_base64():
    public = process_upload(png_bytes(), "shot.png").to_public()
    assert public["has_image"] is True
    assert "image_b64" not in public


# ------------------------------------------------------------------ endpoints
def test_inspect_endpoint(client):
    response = client.post(
        "/v1/attachments/inspect", files={"file": ("notes.txt", b"inspect me", "text/plain")}
    )
    assert response.status_code == 200
    assert response.json()["handled_as"] == "text_read"


def test_capabilities_endpoint_is_honest(client, monkeypatch):
    # Model provisioning is intentionally allowed to activate a model, so do
    # not let another test's hardware-dependent recommendation leak into this
    # assertion about the text-only default.
    from app.core.config import settings

    monkeypatch.setattr(settings, "ollama_model", "qwen3:4b")
    body = client.get("/v1/attachments/capabilities").json()
    assert body["formats"]["text"]["supported"] is True
    assert body["formats"]["documents"]["supported"] is True
    assert body["formats"]["video"]["supported"] is False
    # Default test model is text-only, so vision must be reported as unavailable.
    assert body["formats"]["images"]["supported"] is False


def test_upload_chat_rejects_empty_request(client):
    assert client.post("/v1/chat/upload", data={"message": ""}).status_code == 400
